import csv
import io
from datetime import datetime, timezone
from typing import Any
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .storage import GradeStore

MAX_FILE_BYTES = 5 * 1024 * 1024
ALIASES = {
    "student_id": {"student_id", "id", "學號", "學生編號"},
    "name": {"name", "姓名", "學生姓名"},
    "class_name": {"class", "class_name", "班級", "課程"},
    "coursework": {"coursework", "平時", "平時成績", "作業"},
    "midterm": {"midterm", "期中", "期中成績", "期中考"},
    "final": {"final", "期末", "期末成績", "期末考"},
    "attendance_rate": {"attendance_rate", "出席率", "出席率百分比"},
    "absences": {"absences", "缺席次數", "缺勤"},
}


class ImportService:
    def __init__(self, store: GradeStore) -> None:
        self.store = store

    def preview(self, filename: str, content: bytes) -> dict[str, Any]:
        if len(content) > MAX_FILE_BYTES:
            raise ValueError("檔案不可大於 5 MB")
        suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
        if suffix == "csv":
            rows = self._read_csv(content)
        elif suffix in {"xlsx", "xlsm"}:
            rows = self._read_xlsx(content)
        elif suffix == "xls":
            rows = self._read_xls(content)
        else:
            raise ValueError("僅支援 CSV 或 Excel（.xls、.xlsx）檔案")
        if len(rows) > 500:
            raise ValueError("單次最多匯入 500 位學生")
        known_ids = self.store.student_ids()
        seen_ids: set[str] = set()
        valid = []
        errors = []
        duplicates = 0
        for row_number, raw in enumerate(rows, start=2):
            try:
                record = self._normalize_row(raw)
                if record["student_id"] in seen_ids or record["student_id"] in known_ids:
                    duplicates += 1
                    errors.append({"row": row_number, "student_id": record["student_id"], "message": "學號重複，已略過"})
                    continue
                seen_ids.add(record["student_id"])
                valid.append(record)
            except (TypeError, ValueError) as error:
                errors.append({"row": row_number, "message": str(error)})
        return {
            "filename": filename,
            "total_rows": len(rows),
            "valid_count": len(valid),
            "duplicate_count": duplicates,
            "invalid_count": len(errors) - duplicates,
            "errors": errors,
            "records": valid,
            "preview": valid[:10],
        }

    def commit(self, records: list[dict[str, Any]]) -> dict[str, int]:
        if not records or len(records) > 500:
            raise ValueError("每次匯入需有 1 至 500 筆有效資料")
        inserted = 0
        duplicates = 0
        for record in records:
            student = {**record, "updated_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")}
            if self.store.insert_student(student):
                inserted += 1
            else:
                duplicates += 1
        return {"imported": inserted, "duplicates_skipped": duplicates, "total": self.store.count_students()}

    @staticmethod
    def _read_csv(content: bytes) -> list[dict[str, Any]]:
        text = None
        for encoding in ("utf-8-sig", "cp950", "utf-8"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("無法辨識 CSV 編碼，請另存為 UTF-8 或 Big5")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("檔案沒有標題列")
        return list(reader)

    @staticmethod
    def _read_xlsx(content: bytes) -> list[dict[str, Any]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                raise ValueError("Excel 檔案沒有標題列")
            keys = [str(value).strip() if value is not None else "" for value in headers]
            return [dict(zip(keys, row)) for row in values if any(value is not None for value in row)]
        except (OSError, KeyError, TypeError, BadZipFile, InvalidFileException) as error:
            raise ValueError("Excel 檔案無法讀取，請確認檔案未損毀") from error

    @staticmethod
    def _read_xls(content: bytes) -> list[dict[str, Any]]:
        try:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows == 0:
                raise ValueError("Excel 檔案沒有標題列")
            headers = [str(value).strip() for value in sheet.row_values(0)]
            return [
                dict(zip(headers, sheet.row_values(row_index)))
                for row_index in range(1, sheet.nrows)
                if any(value != "" for value in sheet.row_values(row_index))
            ]
        except (xlrd.XLRDError, IndexError) as error:
            raise ValueError("Excel 檔案無法讀取，請確認檔案未損毀") from error

    @staticmethod
    def _normalize_row(raw: dict[str, Any]) -> dict[str, Any]:
        normalized = {str(key).strip().lower(): value for key, value in raw.items() if key is not None}
        record: dict[str, Any] = {}
        for target, aliases in ALIASES.items():
            value = next((normalized[alias.lower()] for alias in aliases if alias.lower() in normalized), None)
            if value is not None and str(value).strip() != "":
                record[target] = value
        for key in ("student_id", "name", "coursework", "midterm", "final"):
            if key not in record:
                raise ValueError(f"缺少必要欄位或資料：{key}")
        record["student_id"] = str(record["student_id"]).strip()
        record["name"] = str(record["name"]).strip()
        if not record["student_id"] or not record["name"]:
            raise ValueError("學號和姓名不可空白")
        scores = {}
        for component in ("coursework", "midterm", "final"):
            try:
                score = float(record[component])
            except (ValueError, TypeError) as error:
                raise ValueError(f"{component} 必須是數字") from error
            if not 0 <= score <= 100:
                raise ValueError(f"{component} 必須介於 0 到 100 分")
            scores[component] = score
        try:
            attendance_rate = float(record.get("attendance_rate", 100))
            absences = int(float(record.get("absences", 0)))
        except (ValueError, TypeError) as error:
            raise ValueError("出席率或缺席次數格式錯誤") from error
        if not 0 <= attendance_rate <= 100 or absences < 0:
            raise ValueError("出席率需介於 0 到 100，缺席次數不可小於 0")
        return {
            "student_id": record["student_id"],
            "name": record["name"],
            "class_name": str(record.get("class_name", "未分類") or "未分類").strip(),
            "scores": scores,
            "attendance_rate": attendance_rate,
            "absences": absences,
        }